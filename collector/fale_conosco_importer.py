import os
import csv
import io
import hashlib
from datetime import datetime
from typing import Dict, Any, List, Union
import openpyxl

from database.db import upsert_support_ticket

def classify_ticket_category(description: str, analysis: str = "") -> str:
    """Classifica a categoria do chamado com base no relato e análise interna."""
    desc = (description or "").lower()
    ana = (analysis or "").lower()
    combined = f"{desc} {ana}"
    
    if "token" in combined or "0198268" in ana or "totem" in combined or "tolken" in combined:
        return "Erro de Token de Autenticação"
    elif "gov" in combined or "seg.id" in combined or "seg id" in combined or "0198289" in ana or "0198367" in ana or "0198269" in ana:
        return "Login Gov.br / SEG.ID"
    elif any(k in combined for k in ["mudei de escola", "troquei de escola", "transfer", "escola antiga"]) or "0198728" in ana:
        return "Vínculo / Troca de Escola"
    elif any(k in combined for k in ["senha", "esqueci", "id do aluno", "id aluno", "não sei meu id"]):
        return "Recuperação de Senha / ID"
    elif any(k in combined for k in ["reprovad", "nota", "boletim", "falta", "frequência", "frequencia", "presença"]):
        return "Notas e Frequência"
    elif any(k in combined for k in ["compatív", "compativ", "dispositivo", "poco", "redmi", "android 14", "android 12"]):
        return "Compatibilidade de Dispositivo"
    elif any(k in combined for k in ["histórico", "historico", "declaração", "declaracao", "transferência"]):
        return "Documentação / Secretaria"
    elif any(k in combined for k in ["não entra", "nao entra", "tela preta", "trava", "congela", "não abre"]):
        return "Instabilidade / Falha no App"
    return "Outros"

def import_fale_conosco_file(file_path_or_bytes: Union[str, bytes], filename: str = "") -> Dict[str, Any]:
    """Importa chamados do Fale Conosco a partir de Excel (.xlsx) ou CSV."""
    is_xlsx = filename.lower().endswith('.xlsx') or (isinstance(file_path_or_bytes, str) and file_path_or_bytes.lower().endswith('.xlsx'))
    
    rows = []
    if is_xlsx:
        if isinstance(file_path_or_bytes, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
        else:
            wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])
    else:
        # CSV
        text = ""
        if isinstance(file_path_or_bytes, bytes):
            for enc in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
                try:
                    text = file_path_or_bytes.decode(enc)
                    break
                except Exception:
                    continue
        else:
            if os.path.exists(file_path_or_bytes):
                with open(file_path_or_bytes, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read()
            else:
                text = str(file_path_or_bytes)
                
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)

    if not rows:
        return {"success": False, "message": "Arquivo vazio ou inválido."}

    # Mapeamento dinâmico de colunas
    col_idx = {
        'date': 0,
        'user_type': 1,
        'student_name': 2,
        'student_cpf': 3,
        'student_id': 4,
        'birth_date': 5,
        'contact_email': 6,
        'school_name': 7,
        'description': 8,
        'attachment_url': 9,
        'city': 10,
        'status': 11,
        'analysis': 12
    }

    header_found = False
    data_rows = []

    for row_num, row in enumerate(rows):
        if not row or not any(row):
            continue
        first_cell = str(row[0]).strip().lower()
        if 'carimbo' in first_cell or 'data' in first_cell:
            for i, h in enumerate(row):
                hl = str(h).strip().lower()
                if 'carimbo' in hl or 'data' in hl: col_idx['date'] = i
                elif 'você é' in hl or 'perfil' in hl: col_idx['user_type'] = i
                elif 'nome' in hl and 'aluno' in hl: col_idx['student_name'] = i
                elif 'cpf' in hl: col_idx['student_cpf'] = i
                elif 'id' in hl and 'aluno' in hl: col_idx['student_id'] = i
                elif 'nascimento' in hl: col_idx['birth_date'] = i
                elif 'email' in hl or 'contato' in hl: col_idx['contact_email'] = i
                elif 'escola' in hl: col_idx['school_name'] = i
                elif 'problema' in hl or 'descreva' in hl: col_idx['description'] = i
                elif 'anexo' in hl: col_idx['attachment_url'] = i
                elif 'cidade' in hl or 'município' in hl: col_idx['city'] = i
                elif 'status' in hl: col_idx['status'] = i
                elif 'análise' in hl or 'analise' in hl or 'mantis' in hl: col_idx['analysis'] = i
            header_found = True
            continue
            
        data_rows.append(row)

    imported_count = 0
    updated_count = 0

    for row in data_rows:
        def get_val(col_name, default=""):
            idx = col_idx.get(col_name)
            if idx is not None and idx < len(row):
                v = str(row[idx]).strip()
                if v and v.lower() != 'none':
                    return v
            return default

        ticket_date = get_val('date')
        if not ticket_date or 'carimbo' in ticket_date.lower():
            continue

        user_type = get_val('user_type', 'Aluno')
        student_name = get_val('student_name')
        student_cpf = get_val('student_cpf')
        student_id = get_val('student_id')
        birth_date = get_val('birth_date')
        contact_email = get_val('contact_email')
        school_name = get_val('school_name')
        description = get_val('description')
        attachment_url = get_val('attachment_url')
        city = get_val('city')
        status = get_val('status', 'Pendente')
        analysis = get_val('analysis')

        # Formata data se vier com formato datetime do Excel
        if len(ticket_date) > 19 and '.' in ticket_date:
            ticket_date = ticket_date.split('.')[0]

        hash_seed = f"{ticket_date}_{student_cpf}_{student_name}_{description[:30]}"
        ticket_id = "fc_" + hashlib.md5(hash_seed.encode('utf-8')).hexdigest()[:16]
        category = classify_ticket_category(description, analysis)

        ticket = {
            'id': ticket_id,
            'ticket_date': ticket_date,
            'user_type': user_type,
            'student_name': student_name,
            'student_cpf': student_cpf,
            'student_id': student_id,
            'birth_date': birth_date,
            'contact_email': contact_email,
            'school_name': school_name,
            'description': description,
            'attachment_url': attachment_url,
            'city': city,
            'status': status if status else 'Pendente',
            'internal_analysis': analysis,
            'category': category
        }

        is_new = upsert_support_ticket(ticket)
        if is_new:
            imported_count += 1
        else:
            updated_count += 1

    return {
        "success": True,
        "imported": imported_count,
        "updated": updated_count,
        "total": imported_count + updated_count
    }
